"""Lay the platform slides out on the ASHG 2026 speaker template and merge in the organiser's.

`npm run export:slides` writes `platform-slides.json` (the content, from the bundles) and
`captures/` (the charts, from the running app). This script places them on ASHG's own
template, so the deck carries the meeting's title slide, disclosure slides, palette, fonts
and footer, and keeps to ASHG's presenter guidelines: 44 pt titles, a unique title per
slide, nothing projected under 24 pt, and at most six lines of text. Anything longer is in
the speaker notes.

When the organiser's deck is present, every slide in it that this script did not generate
(the literature examples after each module) is carried over after the same module it
followed there. That works on the organiser's hand-merged deck and on this script's own
output, so the organiser can keep editing the built deck and rebuild it after a bundle
change. Filled-in disclosure and title slides in that deck are kept over the template's.

    npm run build:deck -- [--template PATH] [--examples PATH|none] [--out PATH]
"""
import argparse
import copy
import json
import re
import sys
from io import BytesIO
from pathlib import Path

import segno
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE_URL = "https://www.ashg.org/wp-content/uploads/2026/07/2026-PowerPoint-Template_Speakers.pptx"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

WORKSHOP_TITLE = "Modeling Choices in Omics"
WORKSHOP_SUBTITLE = "Implications for Health Disparities in Human Genetics"
AFFILIATION = "Northwestern University Feinberg School of Medicine"
# The workshop organisers, as each wants to be named. Dr. Perera's degrees and title are as
# her lab site gives them (pereralab.mystrikingly.com: Associate Professor, Department of
# Pharmacology and Center for Pharmacogenomics), matching her Feinberg faculty profile.
ORGANIZERS = [
    ("Kynon J. Benjamin, PhD", ["Assistant Professor", "HEART-GeN Lab"]),
    ("Minoli Perera, PharmD, PhD", ["Associate Professor", "Perera Lab"]),
]
FUNDING = "NIH NIMHD R00MD016964 (KJMB)"
# ASHG's logo policy allows institutional logos only on an acknowledgement slide, preferably
# the last, so they appear on the closing slide and nowhere else. Kept out of the repository
# like the template; see README.md for where each comes from. A missing file is skipped.
LOGOS = ["heart-gen.png", "feinberg.png", "nimhd.png"]
# ASHG's own wording for each disclosure, in the plural: neither presenter has anything to
# disclose (confirmed 2026-10-18). An organiser-edited copy in the examples deck wins.
DISCLOSURES = {
    "template:disclosure": "We do not have any relationships to report within the last 24 months with "
                           "ACCME defined ineligible companies.",
    "template:unlabeled": "We will not be discussing unlabeled/investigational uses of medical devices or "
                          "pharmaceuticals during this presentation.",
}

# The template's body region: below the title placeholder, above the footer bars.
BODY_X, BODY_W = 0.92, 11.5
BODY_TOP, BODY_BOTTOM = 2.45, 6.8
MIN_PT = 24  # ASHG: "No text should be smaller than 24 point"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE = "E8E8E8"  # the template's lt2

# Module 6 is about who can afford the compute. These say what to reach for. They are not
# derived from the bundles, so every claim is one the project's own documentation makes;
# the citations are in the notes.
RESOURCE_SLIDES = [
    {
        "kind": "list",
        "tag": "resource:frameworks",
        "eyebrow": "Module 6 · Resources",
        "title": "Same code on CPU or GPU",
        "items": [
            ("PyTorch", "https://pytorch.org", "one code path, CPU or GPU"),
            ("JAX", "https://docs.jax.dev", "compiled for CPU, GPU or TPU"),
            ("CuPy · cuDF", "https://rapids.ai", "NumPy and pandas on the GPU"),
            ("Numba", "https://numba.pydata.org", "Python loops on threads or CUDA"),
            ("Dask", "https://www.dask.org", "a laptop's cores to a cluster"),
        ],
        "notes": (
            "The point for module 6: the permutation scan is infeasible on one node because of hardware, "
            "not because the method is exotic. A genome-wide scan repeats the same regression at millions "
            "of variants, which is exactly the work these libraries spread across cores and GPUs without "
            "changing the statistics.\n\n"
            "PyTorch: tensors and models move between CPU and GPU with .to(device); multi-threaded on CPU. "
            "JAX: NumPy-style code compiled through XLA; vmap turns a one-variant test into an all-variants "
            "test. CuPy is a NumPy-compatible GPU array library; RAPIDS cuDF is the pandas equivalent, and "
            "cudf.pandas runs existing pandas code on the GPU. Numba: @njit(parallel=True) across CPU "
            "threads, @cuda.jit for NVIDIA GPUs. Dask: the same array and dataframe code from one machine to "
            "a cluster, including SLURM through dask-jobqueue.\n\n"
            "https://pytorch.org · https://docs.jax.dev · https://cupy.dev · https://rapids.ai · "
            "https://numba.pydata.org · https://www.dask.org · https://jobqueue.dask.org"
        ),
    },
    {
        "kind": "list",
        "tag": "resource:genetics",
        "eyebrow": "Module 6 · Resources",
        "title": "Genetics tools that already scale",
        "items": [
            ("PLINK 2", "https://www.cog-genomics.org/plink/2.0/", "association and QC, multi-threaded"),
            ("REGENIE", "https://rgcgithub.github.io/regenie/", "whole-genome regression, multi-threaded"),
            ("Hail", "https://hail.is", "Spark clusters; gnomAD, All of Us"),
            ("tensorQTL", "https://github.com/broadinstitute/tensorqtl", "QTL mapping on GPU (PyTorch)"),
            ("Parabricks", "https://docs.nvidia.com/clara/parabricks/latest/", "GPU alignment and variant calling"),
            ("NSF ACCESS", "https://allocations.access-ci.org", "research HPC at no cost (Bridges-2)"),
        ],
        "notes": (
            "PLINK 2: association testing, QC and scoring on biobank-scale genotypes; multi-threaded "
            "(Chang et al., GigaScience 2015). REGENIE: whole-genome regression for quantitative and binary "
            "traits; --threads defaults to all cores but one (Mbatchou et al., Nature Genetics 53:1097-1103, "
            "2021). Hail: genomic dataframes on Apache Spark, from a laptop to a cluster; its README names it "
            "the analysis engine for gnomAD, and the All of Us Researcher Workbench documents Hail on "
            "Dataproc clusters. tensorQTL: cis- and trans-QTL mapping built on PyTorch; the README reports "
            "~200-300x faster than CPU-based implementations, and it also runs on CPU (Taylor-Weiner, Aguet "
            "et al., Genome Biology 2019). NVIDIA Parabricks: GPU versions of BWA-MEM alignment, DeepVariant "
            "and other callers; free public container, paid enterprise support optional. NSF ACCESS: "
            "allocations for research and education, including Bridges-2 at PSC, the platform module 6 "
            "prices; cloud providers also run research credit programs."
        ),
    },
]


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--manifest", type=Path, default=HERE / "platform-slides.json")
    parser.add_argument("--template", type=Path, default=HERE / "2026-PowerPoint-Template_Speakers.pptx")
    parser.add_argument("--examples", default=str(HERE / "ASHG_KJ_2026_platform-slides.pptx"),
                        help="the organiser's deck to take example slides from, or 'none'")
    parser.add_argument("--logos", type=Path, default=HERE / "logos")
    parser.add_argument("--out", type=Path, default=HERE / "ASHG2026-modeling-choices.pptx")
    return parser.parse_args()


# --- Slide plumbing -----------------------------------------------------------------------

def tag(slide):
    return slide._element.cSld.get("name", "")


def set_tag(slide, value):
    slide._element.cSld.set("name", value)


def slide_text(slide):
    return "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


def remove_slide(prs, slide):
    id_list = prs.slides._sldIdLst
    for entry in list(id_list):
        if prs.part.related_part(entry.rId) is slide.part:
            prs.part.drop_rel(entry.rId)
            id_list.remove(entry)
            return
    raise ValueError("slide not in presentation")


def layout(prs, name):
    return next(item for item in prs.slide_layouts if item.name == name)


def set_notes(slide, text):
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        raise ValueError("the template's notes master has no body placeholder")
    frame.text = text


def materialise_geometry(slide):
    """Give inherited placeholders explicit positions so they survive a change of layout."""
    for shape in slide.placeholders:
        if shape._element.spPr.find(qn("a:xfrm")) is None:
            try:
                left, top, width, height = shape.left, shape.top, shape.width, shape.height
            except AttributeError:
                continue
            if None not in (left, top, width, height):
                shape.left, shape.top, shape.width, shape.height = left, top, width, height


def materialise_text_style(slide):
    """Write each placeholder's inherited list style onto the shape itself.

    A placeholder's type size, font and colour usually come from its layout and master. On
    the template's layouts those are 28-44 pt Poppins, so a 14 pt bullet list arrives at
    28 pt and runs off the slide. Each level takes the nearest definition: the shape's own,
    then its layout placeholder's, then its master placeholder's, then the master's text style.
    """
    tx_styles = slide.slide_layout.slide_master._element.find(qn("p:txStyles"))
    for shape in slide.placeholders:
        body = shape._element.find(qn("p:txBody"))
        if body is None:
            continue
        kind = shape.placeholder_format.type
        is_title = kind in (1, 3)  # TITLE, CENTER_TITLE
        sources = [body.find(qn("a:lstStyle"))]
        inherited = shape
        for _ in range(2):  # layout placeholder, then master placeholder
            inherited = getattr(inherited, "_base_placeholder", None)
            if inherited is None:
                break
            inherited_body = inherited._element.find(qn("p:txBody"))
            sources.append(inherited_body.find(qn("a:lstStyle")) if inherited_body is not None else None)
        if tx_styles is not None:
            sources.append(tx_styles.find(qn("p:titleStyle" if is_title else "p:bodyStyle")))
        levels = []
        for level in range(1, 10):
            tag = qn(f"a:lvl{level}pPr")
            found = next((source.find(tag) for source in sources if source is not None and source.find(tag) is not None), None)
            if found is not None:
                levels.append(copy.deepcopy(found))
        if not levels:
            continue
        existing = body.find(qn("a:lstStyle"))
        style = body.makeelement(qn("a:lstStyle"), {})
        style.extend(levels)
        if existing is not None:
            body.replace(existing, style)
        else:
            body.insert(1, style)  # after bodyPr


def copy_shapes(source, target, skip=None):
    """Copy every shape on `source` onto `target`, re-pointing images and links at `target`."""
    tree = target.shapes._spTree
    for child in source.shapes._spTree:
        if child.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")) or child is skip:
            continue
        clone = copy.deepcopy(child)
        for element in clone.iter():
            for attribute, rId in list(element.attrib.items()):
                if not attribute.startswith(f"{{{R_NS}}}"):
                    continue
                rel = source.part.rels[rId]
                if rel.is_external:
                    new_rId = target.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                elif rel.reltype == RT.IMAGE:
                    _, new_rId = target.part.get_or_add_image_part(BytesIO(rel.target_part.blob))
                else:
                    raise ValueError(f"slide {source.slide_id}: cannot carry a {rel.reltype} relationship")
                element.set(attribute, new_rId)
        tree.append(clone)
    background = source._element.cSld.bg
    if background is not None:
        target._element.cSld.insert(0, copy.deepcopy(background))
    if source.has_notes_slide and source.notes_slide.notes_text_frame is not None:
        notes = source.notes_slide.notes_text_frame.text
        if notes.strip():
            set_notes(target, notes)


def clone_slide(source, prs, layout_name):
    materialise_geometry(source)
    materialise_text_style(source)
    slide = prs.slides.add_slide(layout(prs, layout_name))
    for placeholder in list(slide.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    copy_shapes(source, slide)
    return slide


def clone_example(source, prs):
    """An example slide, with its title (if it has one) in the template's title placeholder.

    The organiser's titles sat wherever their own layouts put them: on top of a figure, or
    half off the slide. Moved to the template's title position, they read like every other
    slide's, and the rest of the slide is fitted into the body region below.
    """
    title = source.shapes.title
    if title is None or not title.text_frame.text.strip():
        slide = clone_slide(source, prs, "Blank")
        # Kept above the template's footer bars, where several citations had landed.
        fit_into(slide, list(slide.shapes), box=(0.3, 0.25, 12.73, 6.6))
        return slide
    materialise_geometry(source)
    materialise_text_style(source)
    slide = prs.slides.add_slide(layout(prs, "Title and Content"))
    for placeholder in list(slide.placeholders):
        if placeholder.placeholder_format.type != 1:
            placeholder._element.getparent().remove(placeholder._element)
    slide.shapes.title.text = " ".join(title.text_frame.text.split())
    copy_shapes(source, slide, skip=title._element)
    fit_into(slide, [shape for shape in slide.shapes
                     if not (shape.is_placeholder and shape.placeholder_format.type == 1)])
    return slide


def fit_into(slide, shapes, box=(0.6, 1.95, 12.13, 4.9)):
    """Scale and move `shapes` together so their bounding box sits inside `box` (inches)."""
    placed = [shape for shape in shapes if None not in (shape.left, shape.top, shape.width, shape.height)]
    if not placed:
        return
    left = min(shape.left for shape in placed)
    top = min(shape.top for shape in placed)
    right = max(shape.left + shape.width for shape in placed)
    bottom = max(shape.top + shape.height for shape in placed)
    bx, by, bw, bh = (Inches(value) for value in box)
    if left >= bx and top >= by and right <= bx + bw and bottom <= by + bh:
        return
    scale = min(1.0, bw / (right - left), bh / (bottom - top))
    offset_x = bx + (bw - (right - left) * scale) / 2
    for shape in placed:
        shape.left = int(offset_x + (shape.left - left) * scale)
        shape.top = int(by + (shape.top - top) * scale)
        shape.width = int(shape.width * scale)
        shape.height = int(shape.height * scale)
        if shape.has_text_frame and scale < 1:
            scale_text(shape.text_frame, scale)


def scale_text(frame, scale):
    """Shrink the type with its box; a smaller box with the same type overflows the slide."""
    body = frame._txBody.bodyPr
    current = 1.0
    for child in list(body):
        if child.tag in (qn("a:noAutofit"), qn("a:spAutoFit"), qn("a:normAutofit")):
            current = int(child.get("fontScale", "100000")) / 100000
            body.remove(child)
    # bodyPr's first children are prstTxWarp and the autofit choice; put it first.
    body.insert(0, body.makeelement(qn("a:normAutofit"), {"fontScale": str(round(current * scale * 100000))}))


def fill_disclosure(slide, statement):
    """Replace the template's instructions, both options and the company table with one statement."""
    for shape in list(slide.shapes):
        if shape.has_table:
            shape._element.getparent().remove(shape._element)
    body = next(shape for shape in slide.placeholders if shape.placeholder_format.type != 1)
    frame = body.text_frame
    frame.text = ""
    style_run(frame.paragraphs[0].add_run(), statement, 32)


def downscale_images(prs, max_width=2560, threshold=2_500_000):
    """The template's photographs are 10-20 MB each; the projector is 1280x720."""
    for part in prs.part.package.iter_parts():
        if part.content_type not in ("image/jpeg", "image/png") or len(part.blob) < threshold:
            continue
        image = Image.open(BytesIO(part.blob))
        if image.width <= max_width and part.content_type != "image/jpeg":
            continue
        if image.width > max_width:
            image = image.resize((max_width, round(image.height * max_width / image.width)), Image.LANCZOS)
        buffer = BytesIO()
        if part.content_type == "image/jpeg":
            image.convert("RGB").save(buffer, "JPEG", quality=85, optimize=True)
        else:
            image.save(buffer, image.format or "PNG", optimize=True)
        if len(buffer.getvalue()) < len(part.blob):
            part._blob = buffer.getvalue()


# --- Text helpers -------------------------------------------------------------------------

def style_run(run, text, size, bold=False, theme=None, rgb=None, link=None, caps=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    if link:
        run.hyperlink.address = link
        theme = theme or MSO_THEME_COLOR.HYPERLINK
    if rgb is not None:
        run.font.color.rgb = rgb
    elif theme is not None:
        run.font.color.theme_color = theme
    if caps:
        run.font._rPr.set("cap", "all")


def textbox(slide, x, y, w, h, paragraphs, anchor=MSO_ANCHOR.TOP, space_after=0):
    """paragraphs: a list of paragraphs, each a list of (text, style) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, runs in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(space_after)
        for text, style in runs:
            assert style.get("size", MIN_PT) >= MIN_PT, f"{text!r}: under {MIN_PT} pt"
            style_run(paragraph.add_run(), text, **style)
    return box


def content_slide(prs, spec, tag_value):
    """The template's 'Title and Content' slide: title placeholder, our own body shapes."""
    slide = prs.slides.add_slide(layout(prs, "Title and Content"))
    set_tag(slide, tag_value)
    for placeholder in list(slide.placeholders):
        if placeholder.placeholder_format.type != 1:  # keep the title, drop the content box
            placeholder._element.getparent().remove(placeholder._element)
    slide.shapes.title.text = spec["title"]
    textbox(slide, BODY_X, 1.85, BODY_W, 0.45,
            [[(spec["eyebrow"], {"size": 24, "bold": True, "theme": MSO_THEME_COLOR.ACCENT_4, "caps": True})]])
    if spec.get("notes"):
        set_notes(slide, spec["notes"])
    return slide


def table(slide, header, rows, widths, top=BODY_TOP):
    shape = slide.shapes.add_table(len(rows) + (1 if header else 0), len(widths), Inches(BODY_X), Inches(top),
                                   Inches(sum(widths)), Inches(0.5 * (len(rows) + 1)))
    grid = shape.table
    # "No Style, No Grid": the fills and rules below are the whole look.
    grid._tbl.tblPr.find(qn("a:tableStyleId")).text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    grid.first_row = grid.horz_banding = False
    for index, width in enumerate(widths):
        grid.columns[index].width = Inches(width)
    body = ([header] if header else []) + rows
    for r, values in enumerate(body):
        grid.rows[r].height = Inches(0.5)
        is_header = bool(header) and r == 0
        for c, value in enumerate(values):
            cell = grid.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.06)
            frame = cell.text_frame
            frame.word_wrap = True
            style = {"size": MIN_PT, "bold": is_header or c == 0}
            if is_header:
                style["rgb"] = WHITE
            style_run(frame.paragraphs[0].add_run(), value, **style)
            if is_header:
                cell.fill.solid()
                cell.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
            else:
                cell.fill.background()
            tcPr = cell._tc.get_or_add_tcPr()
            # The schema puts the borders before the cell fill; PowerPoint repairs the file otherwise.
            for position, edge in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                line = tcPr.makeelement(qn(edge), {"w": str(Pt(1)), "cmpd": "sng"})
                fill = line.makeelement(qn("a:solidFill"), {})
                fill.append(fill.makeelement(qn("a:srgbClr"), {"val": RULE}))
                line.append(fill)
                tcPr.insert(position, line)
    return shape


# --- Slide kinds --------------------------------------------------------------------------

def build_join(prs, spec, tag_value):
    slide = content_slide(prs, spec, tag_value)
    # Text on the left, the QR code on the right: a phone camera is faster than typing.
    text_w, qr_size = 7.7, 3.3
    qr_x = BODY_X + BODY_W - qr_size
    textbox(slide, BODY_X, BODY_TOP + 0.1, text_w, 0.9,
            [[(spec["displayUrl"], {"size": 36, "bold": True, "link": spec["displayHref"],
                                    "theme": MSO_THEME_COLOR.ACCENT_2})]])
    no_install, poll, question = spec["lines"]
    textbox(slide, BODY_X, BODY_TOP + 1.2, text_w, 3.0, [
        [(no_install, {"size": 28})],
        [(poll, {"size": 24, "theme": MSO_THEME_COLOR.ACCENT_1})],
        [(question, {"size": 28, "bold": True})],
    ], space_after=14)
    image = BytesIO()
    # Error correction M survives a projector's glare; the 4-module quiet zone is the spec's.
    segno.make(spec["url"], error="m").save(image, kind="png", scale=20, border=4)
    image.seek(0)
    slide.shapes.add_picture(image, Inches(qr_x), Inches(BODY_TOP - 0.15), Inches(qr_size), Inches(qr_size))
    textbox(slide, qr_x, BODY_TOP - 0.15 + qr_size, qr_size, 0.45,
            [[("Scan to join", {"size": 24, "bold": True})]]).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER


def build_table(prs, spec, tag_value):
    slide = content_slide(prs, spec, tag_value)
    widths = [7.9, 3.6] if spec["header"][1] == "Time" else [4.6, 6.9]
    table(slide, spec["header"], spec["rows"], widths)


def build_steps(prs, spec, tag_value):
    slide = content_slide(prs, spec, tag_value)
    column, gap = 2.6, (BODY_W - 4 * 2.6) / 3
    for position, (name, short) in enumerate(spec["steps"]):
        x = BODY_X + position * (column + gap)
        textbox(slide, x, BODY_TOP + 0.1, column, 0.9,
                [[(str(position + 1), {"size": 54, "bold": True, "theme": MSO_THEME_COLOR.ACCENT_1})]])
        textbox(slide, x, BODY_TOP + 1.1, column, 0.6, [[(name, {"size": 28, "bold": True})]])
        textbox(slide, x, BODY_TOP + 1.75, column, 2.0, [[(short, {"size": 24})]])


def build_commit(prs, spec, tag_value):
    slide = content_slide(prs, spec, tag_value)
    textbox(slide, BODY_X, BODY_TOP, BODY_W, BODY_BOTTOM - BODY_TOP,
            [[(spec["question"], {"size": 24, "bold": True})]]
            + [[(option, {"size": 24})] for option in spec["options"]], space_after=8)


def build_answer(prs, spec, tag_value):
    slide = content_slide(prs, spec, tag_value)
    textbox(slide, BODY_X, BODY_TOP, BODY_W, BODY_BOTTOM - BODY_TOP, [
        [("Answer", {"size": 28, "theme": MSO_THEME_COLOR.ACCENT_2, "bold": True})],
        [(spec["answer"], {"size": 40, "bold": True, "theme": MSO_THEME_COLOR.ACCENT_1})],
    ], anchor=MSO_ANCHOR.MIDDLE, space_after=12)


def build_chart(prs, spec, tag_value, manifest_dir):
    slide = content_slide(prs, spec, tag_value)
    path = manifest_dir / spec["image"]
    with Image.open(path) as image:
        width, height = image.size
    box_w, box_h = BODY_W, BODY_BOTTOM - BODY_TOP + 0.05
    scale = min(box_w / width, box_h / height)
    w, h = width * scale, height * scale
    slide.shapes.add_picture(str(path), Inches(BODY_X + (box_w - w) / 2), Inches(BODY_TOP), Inches(w), Inches(h))


def build_list(prs, spec, tag_value):
    slide = content_slide(prs, spec, tag_value)
    assert len(spec["items"]) <= 6, f"{spec['title']}: more than six lines"
    textbox(slide, BODY_X, BODY_TOP, BODY_W, BODY_BOTTOM - BODY_TOP, [
        [(name, {"size": 28, "bold": True, "link": link}), (f"  {description}", {"size": 28})]
        for name, link, description in spec["items"]
    ], space_after=10)


def build_closing(prs, logo_dir):
    """Thank-you slide: the organisers, the funding, and the only logos in the deck."""
    spec = {"eyebrow": "Workshop organizers", "title": "Thank you", "notes": ""}
    slide = content_slide(prs, spec, "closing:thanks")
    column = BODY_W / len(ORGANIZERS)
    for position, (name, roles) in enumerate(ORGANIZERS):
        textbox(slide, BODY_X + position * column, BODY_TOP, column - 0.3, 1.5,
                [[(name, {"size": 28, "bold": True})]] + [[(role, {"size": 24})] for role in roles])
    textbox(slide, BODY_X, BODY_TOP + 1.45, BODY_W, 1.0, [
        [(AFFILIATION, {"size": 24, "bold": True, "theme": MSO_THEME_COLOR.ACCENT_1})],
        [("Funding: ", {"size": 24, "bold": True}), (FUNDING, {"size": 24})],
    ], space_after=6)
    logos = [logo_dir / name for name in LOGOS if (logo_dir / name).exists()]
    missing = [name for name in LOGOS if not (logo_dir / name).exists()]
    if missing:
        print(f"Closing slide: no {', '.join(missing)} in {logo_dir}; left off.")
    # One row above the footer bars, every logo the same height, the row centred.
    height, gap, top = 0.75, 0.6, 5.85
    widths = []
    for path in logos:
        with Image.open(path) as image:
            widths.append(height * image.width / image.height)
    x = BODY_X + (BODY_W - sum(widths) - gap * (len(widths) - 1)) / 2
    for path, width in zip(logos, widths):
        slide.shapes.add_picture(str(path), Inches(x), Inches(top), Inches(width), Inches(height))
        x += width + gap
    return slide


BUILDERS = {
    "join": build_join, "table": build_table, "steps": build_steps, "commit": build_commit,
    "answer": build_answer, "list": build_list,
}


# --- Organiser deck -----------------------------------------------------------------------

MODULE_EYEBROW = re.compile(r"^Module (\d+) ·")


def is_generated(slide):
    if tag(slide).startswith(("platform:", "resource:", "template:", "closing:")):
        return True
    # The hand-merged deck predates tags: its platform slides are pptxgenjs's DEFAULT layout
    # with the workshop URL on every one.
    return slide.slide_layout.name == "DEFAULT" and "cloudfront.net" in slide_text(slide)


def read_examples(path):
    """Example slides grouped by the module they follow (0 = before module 1)."""
    deck = Presentation(path)
    buckets, overrides, module = {}, {}, 0
    for slide in deck.slides:
        if tag(slide).startswith("template:"):
            overrides[tag(slide)] = slide
        if is_generated(slide):
            match = re.match(r"platform:m(\d+):", tag(slide))
            first_text = next((s.text_frame.text for s in slide.shapes if s.has_text_frame), "")
            legacy = MODULE_EYEBROW.match(first_text)
            if match or legacy:
                module = int((match or legacy).group(1))
            continue
        buckets.setdefault(module, []).append(slide)
    return deck, buckets, overrides


# --- Assembly -----------------------------------------------------------------------------

def main():
    args = parse_args()
    if not args.template.exists():
        sys.exit(f"ASHG template not found at {args.template}.\nDownload it from {TEMPLATE_URL}")
    manifest = json.loads(args.manifest.read_text())
    prs = Presentation(args.template)

    examples_path = None if args.examples == "none" else Path(args.examples)
    buckets, overrides = {}, {}
    if examples_path is not None and examples_path.exists():
        _, buckets, overrides = read_examples(examples_path)
    elif examples_path is not None:
        print(f"No organiser deck at {examples_path}; building platform slides only.")

    # Keep the template's title and two disclosure slides; drop its samples.
    keep = {}
    for slide in list(prs.slides):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        if slide.slide_layout.name == "Title Slide" and "template:title" not in keep:
            keep["template:title"] = slide
        elif title == "Financial Disclosure":
            keep["template:disclosure"] = slide
        elif title == "Unlabeled/Investigational Uses":
            keep["template:unlabeled"] = slide
        else:
            remove_slide(prs, slide)
    assert set(keep) == {"template:title", "template:disclosure", "template:unlabeled"}, keep.keys()
    for name, slide in keep.items():
        set_tag(slide, name)
    keep["template:title"].shapes.title.text = WORKSHOP_TITLE
    # The template's boxes are 7.4 in wide, sized for a short title; widened under the meeting
    # logo, the title fits on one line and the subtitle's three lines stay on the slide.
    title_slide = keep["template:title"]
    title_box, subtitle_box = title_slide.shapes.title, title_slide.placeholders[1]
    title_box.left, title_box.top, title_box.width, title_box.height = Inches(0.63), Inches(2.6), Inches(11.8), Inches(1.4)
    subtitle_box.left, subtitle_box.top = Inches(0.63), Inches(4.2)
    subtitle_box.width, subtitle_box.height = Inches(11.8), Inches(2.6)
    subtitle = subtitle_box.text_frame
    subtitle.text = ""
    lines = [(WORKSHOP_SUBTITLE, False), (" · ".join(name for name, _ in ORGANIZERS), True),
             (AFFILIATION, False)]
    for index, (text, bold) in enumerate(lines):
        paragraph = subtitle.paragraphs[0] if index == 0 else subtitle.add_paragraph()
        style_run(paragraph.add_run(), text, MIN_PT, bold=bold, rgb=WHITE)
    for name, statement in DISCLOSURES.items():
        fill_disclosure(keep[name], statement)
    # Disclosures the organiser has filled in replace the template's instructions.
    for name, source in overrides.items():
        if name in keep:
            replacement = clone_slide(source, prs, source.slide_layout.name)
            set_tag(replacement, name)
            remove_slide(prs, keep[name])
            keep[name] = replacement

    manifest_dir = args.manifest.parent
    order = [keep["template:title"], keep["template:disclosure"], keep["template:unlabeled"]]
    seen_examples = 0

    def add_examples(module):
        nonlocal seen_examples
        for source in buckets.get(module, []):
            order.append(clone_example(source, prs))
            seen_examples += 1

    # Each module's examples follow its last slide; the resource slides follow the last
    # module's, which is the scalability module.
    last_slide = {spec["module"]: position for position, spec in enumerate(manifest["slides"]) if spec.get("module")}
    for position, spec in enumerate(manifest["slides"]):
        module = spec.get("module")
        tag_value = f"platform:m{module}:{position}" if module else f"platform:{spec['kind']}:{position}"
        if spec["kind"] == "chart":
            build_chart(prs, spec, tag_value, manifest_dir)
        else:
            BUILDERS[spec["kind"]](prs, spec, tag_value)
        order.append(prs.slides[-1])
        if spec["kind"] == "steps":
            add_examples(0)
        if module and last_slide[module] == position:
            add_examples(module)
            if module == max(last_slide):
                for resource in RESOURCE_SLIDES:
                    build_list(prs, resource, resource["tag"])
                    order.append(prs.slides[-1])
    for module in sorted(set(buckets) - set(last_slide) - {0}):
        add_examples(module)
    order.append(build_closing(prs, args.logos))

    # Put the slides in the order built above.
    id_list = prs.slides._sldIdLst
    by_part = {prs.part.related_part(entry.rId): entry for entry in id_list}
    for entry in list(id_list):
        id_list.remove(entry)
    for slide in order:
        id_list.append(by_part[slide.part])
    assert len(id_list) == len(by_part), "a built slide was left out of the order"

    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title is not None and tag(s).startswith(("platform:", "resource:"))]
    repeated = sorted({title for title in titles if titles.count(title) > 1})
    assert not repeated, f"slide titles are not unique: {repeated}"

    # python-pptx names a new slide after the current slide count, so a slide added after one
    # was removed can take a live slide's part name, and one overwrites the other on save.
    for number, slide in enumerate(prs.slides, start=1):
        slide.part.partname = PackURI(f"/ppt/slides/slide{number}.xml")

    downscale_images(prs)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.out)
    size = args.out.stat().st_size / 1e6
    print(f"Wrote {args.out} ({len(prs.slides)} slides, {seen_examples} from the organiser deck, {size:.0f} MB)")


if __name__ == "__main__":
    main()
